"""Reading a State Street (SSGA) daily holdings workbook, the one layout all its funds share.

A few header lines (``Fund Name:``, ``Ticker Symbol:``, ``Holdings: As of 01-Oct-2026``), a
table whose header row starts with ``Name`` (equity funds: Name, Ticker, Identifier, SEDOL,
Weight, Sector, Shares Held, Local Currency; bond funds have no Ticker and add Coupon, Par
Value, Market Value, Maturity), then disclaimer text. Weights are percent.
"""

import io
import re
from dataclasses import dataclass
from datetime import date, datetime

import pandas as pd

_AS_OF = re.compile(r"As of (\d{2}-[A-Za-z]{3}-\d{4})")


@dataclass(frozen=True)
class Workbook:
    fund: str | None  # the ticker the sheet says it is for
    as_of: date | None
    table: pd.DataFrame  # one row per line, columns as named in the sheet


def read_workbook(payload: bytes) -> Workbook:
    """Parse the sheet; the table ends at the first line with only a first cell (the disclaimers).
    A line without a weight stays (the caller drops it): a blank cell must not cut the table."""
    import openpyxl  # noqa: PLC0415 - algotrade-sources-only dependency, loaded where needed

    sheet = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True).active
    rows = [r for r in sheet.iter_rows(values_only=True) if any(v is not None for v in r)]
    as_of, fund = None, None
    for row in rows[:6]:
        if row[0] == "Ticker Symbol:" and row[1]:
            fund = str(row[1]).strip().upper()
        match = _AS_OF.search(" ".join(str(v) for v in row if v))
        if match:
            as_of = datetime.strptime(match.group(1), "%d-%b-%Y").date()  # noqa: DTZ007
    header = next((i for i, r in enumerate(rows) if r[0] == "Name"), None)
    if header is None:
        raise ValueError("the workbook has no table (no row starting with Name)")
    # Columns by their position in the raw row (an empty header cell does not shift the others).
    columns = {str(v): i for i, v in enumerate(rows[header]) if v is not None}
    if "Weight" not in columns:  # a changed layout is a parse failure, never a guess
        raise ValueError(f"the workbook's table has no Weight column (columns: {list(columns)})")
    body = []
    for row in rows[header + 1 :]:
        if all(v is None for v in row[1:]):  # disclaimer text after the table
            break
        body.append([row[i] if i < len(row) else None for i in columns.values()])
    return Workbook(fund, as_of, pd.DataFrame(body, columns=list(columns)))
