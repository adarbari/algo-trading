"""etf-holdings review fixes: what a bad read may not replace, false links, ranking, counts."""

import io
from collections.abc import Callable, Mapping
from datetime import date, timedelta
from typing import Any

import openpyxl
import pandas as pd

from algotrade.data.resolver import SymbolResolver
from algotrade.storage.runs import RunRecord, RunStatus
from algotrade.storage.tables.writers import StoreWriter
from algotrade_ingestion.tasks.market.etf_holdings import (
    HoldingsSources,
    Previous,
    Problem,
    fund_frame,
    positions,
    sanity_problem,
)
from algotrade_sources.framework.holdings import holdings_frame
from tests.apps.ingestion.tasks.market.test_etf_holdings import (
    DAY,
    SSGA,
    holdings,
    ishares,
    nport,
    run,
    ssga,
    world,
)

LATER = DAY + timedelta(days=8)
Rows = list[tuple[Any, ...]]


def xlk_file(edit: Callable[[Rows], Rows]) -> bytes:
    """The recorded XLK workbook with its rows (header lines, table, disclaimer) changed."""
    sheet = openpyxl.load_workbook(SSGA / "holdings-daily-us-en-xlk.xlsx").active
    rows = edit([tuple(r) for r in sheet.iter_rows(values_only=True)])
    book = openpyxl.Workbook()
    for row in rows:
        book.active.append(list(row))
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def first_read_then(files: Mapping[str, bytes]) -> tuple[StoreWriter, RunRecord, RunRecord]:
    """A full read of every fund, then (a week later) a forced one where XLK's file is ``files``."""
    writer, sources = world()
    first = run(writer, sources)
    again = HoldingsSources([ssga(files=files), ishares(), nport()], 7, 100, "optionable", False)
    return writer, first, run(writer, again, LATER)


def test_a_truncated_file_that_still_parses_keeps_last_reads_rows() -> None:
    def cut(rows: Rows) -> Rows:
        header = next(i for i, r in enumerate(rows) if r[0] == "Name")
        return [*rows[: header + 4], rows[-1]]  # three holdings and the disclaimer

    writer, first, second = first_read_then({"xlk.xlsx": xlk_file(cut)})
    assert first.items["XLK"].startswith("OK") and second.status is RunStatus.PARTIAL
    assert second.items["XLK"].startswith("FAILED: 3 positions, was 12")
    assert second.stats["rejected"] == 1
    assert len(holdings(writer, "XLK", LATER)) == 12  # last week's rows stay


def test_a_workbook_without_a_weight_column_is_a_failed_read_not_a_snapshot() -> None:
    def rename(rows: Rows) -> Rows:
        return [tuple("Wt" if v == "Weight" else v for v in r) for r in rows]

    writer, _, second = first_read_then({"xlk.xlsx": xlk_file(rename)})
    assert second.items["XLK"].startswith("FETCH_ERROR")
    assert "no Weight column" in second.items["XLK"]
    kept = holdings(writer, "XLK", LATER)
    assert len(kept) == 12 and set(kept["as_of"]) == {date(2026, 10, 1)}


def test_an_older_file_does_not_replace_a_newer_read() -> None:
    def back(rows: Rows) -> Rows:
        return [
            tuple(v.replace("01-Oct-2026", "24-Sep-2026") if isinstance(v, str) else v for v in r)
            for r in rows
        ]

    writer, _, second = first_read_then({"xlk.xlsx": xlk_file(back)})
    assert second.items["XLK"].startswith("FAILED: as of 2026-09-24, older than the stored")
    assert set(holdings(writer, "XLK", LATER)["as_of"]) == {date(2026, 10, 1)}


def test_weights_far_from_100_percent_are_rejected_except_for_geared_and_n_port_funds() -> None:
    writer, base = world()
    checked = HoldingsSources(base.issuers, 7, 100, "optionable", True)
    record = run(writer, checked)  # the recorded files are trimmed: XLK's lines add up to ~75%
    assert record.items["XLK"].startswith("FAILED: weights add up to")
    assert holdings(writer, "XLK").empty
    assert record.items["VTI"].startswith("OK")  # N-PORT: not checked
    geared_writer, geared_base = world(geared=("XLK",))
    geared = HoldingsSources(geared_base.issuers, 7, 100, "optionable", True)
    assert run(geared_writer, geared).items["XLK"].startswith("OK")  # swaps: no 100% rule


def test_an_n_port_report_is_not_shown_before_it_was_filed() -> None:
    """VTI's period-end 2026-06-30 report was filed on 2026-08-28. Run for an earlier session,
    the rows are stored under it, but a read for a day before the filing must not see them."""
    writer, sources = world()
    early = date(2026, 8, 1)
    record = run(writer, sources, early, only=("VTI",))
    assert record.items["VTI"].startswith("OK")
    stored = holdings(writer, "VTI", date(2026, 8, 28))
    assert set(stored["filed"]) == {date(2026, 8, 28)} and set(stored["as_of"]) == {
        date(2026, 6, 30)
    }
    assert holdings(writer, "VTI", date(2026, 8, 27)).empty  # public on the 28th, not before


def problem_text(problem: Problem | None) -> str:
    return "" if problem is None else problem.text


def frame(*lines: tuple[str, float, str]) -> pd.DataFrame:
    return holdings_frame(
        {"holding_name": n, "weight": w, "asset_class": k, "us_listed": False} for n, w, k in lines
    )


def stored(lines: pd.DataFrame) -> pd.DataFrame:
    return lines.assign(identifier=None, filed=None, holding_symbol=None, sector=None, shares=None)


def test_sanity_checks_in_isolation() -> None:
    both = frame(("A", 0.6, "Equity"), ("B", 0.4, "Equity"))
    one = frame(("A", 0.7, "Equity"))
    assert sanity_problem(both, DAY, None, False, True) is None
    assert "add up to 70.0%" in problem_text(sanity_problem(one, DAY, None, False, True))
    assert sanity_problem(one, DAY, None, True, True) is None  # geared
    assert sanity_problem(one, DAY, None, False, False) is None  # N-PORT
    assert "2 positions, was 100" in problem_text(
        sanity_problem(both, DAY, Previous(DAY, 100), False, True)
    )
    assert sanity_problem(both, DAY, Previous(DAY, 3), False, True) is None  # small funds vary
    older = sanity_problem(both, date(2026, 9, 1), Previous(DAY, 2), False, True)
    assert "older than" in problem_text(older)
    assert sanity_problem(both, DAY, Previous(DAY, 2), False, True) is None  # same date, same size


def test_positions_do_not_count_cash_futures_or_fx_lines() -> None:
    lines = frame(
        ("A", 0.5, "Equity"), ("Cash", 0.3, "Cash"), ("ES", 0.1, "Futures"), ("EUR", 0.1, "FX")
    )
    assert positions(lines) == 1
    rows = fund_frame("EQ:F", "F", DAY, stored(lines), 100, SymbolResolver(), {})
    assert set(rows["holdings_count"]) == {1} and len(rows) == 4  # all four lines are stored


def test_inverse_funds_keep_their_big_negative_lines_in_the_top() -> None:
    lines = frame(
        ("T-bill", 0.2, "Fixed Income"),
        ("Swap", -0.9, "Derivative"),
        ("Swap 2", -0.8, "Derivative"),
    )
    assert list(lines["holding_name"]) == ["Swap", "Swap 2", "T-bill"]
    rows = fund_frame("EQ:S", "S", DAY, stored(lines), 2, SymbolResolver(), {})
    assert list(rows["weight"]) == [-0.9, -0.8]  # the sign is kept, the T-bill falls out


def test_a_cusip_printed_for_a_foreign_line_never_links_a_later_line() -> None:
    """Telus (CUSIP 87971M103) trades in Canada as T; the ticker T is AT&T in the universe."""
    resolver = SymbolResolver(ids={"T": "EQ:ATT"}, symbols={"EQ:ATT": "T"})
    lines = stored(frame(("Telus Corp", 0.5, "Equity"), ("A bond", 0.5, "Fixed Income")))
    lines["identifier"] = "87971M103"
    unmapped = fund_frame(
        "EQ:F", "F", DAY, lines, 100, resolver, {}
    )  # a foreign line never fed the map
    assert unmapped["holding_id"].isna().all() and unmapped["holding_symbol"].isna().all()
    mapped = fund_frame("EQ:F", "F", DAY, lines, 100, resolver, {"87971M103": ("T", "EQ:ATT")})
    equity = mapped[mapped["asset_class"] == "Equity"].iloc[0]
    bond = mapped[mapped["asset_class"] == "Fixed Income"].iloc[0]
    assert equity["holding_id"] == "EQ:ATT" and pd.isna(bond["holding_id"])  # equity lines only
